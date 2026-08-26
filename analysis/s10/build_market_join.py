#!/usr/bin/env python3
"""Join the study's occupations to the established market-fitted measures.

Per Matthew (2026-08-26): the convergent/discriminant comparisons against
the classical scales, categorical throughout — continuous anchors are cut
to quartiles downstream. No member placement is needed anywhere: every
anchor attaches at the occupation level, and members inherit through the
study's existing double-blind O*NET-SOC coding (401 distinct codes, 5,188
members). Sources and fetch provenance: market_anchors/PROVENANCE.md.

Anchors joined per O*NET-SOC code:
  isei      ISEI-08 (Ganzeboom), via BLS 2010-SOC x ISCO-08 crosswalk;
            where one SOC maps to several ISCO-08 codes, the mean of the
            mapped scores is taken and the multiplicity recorded
  siops     SIOPS/Treiman-08, same path
  wage      OES May-2024 national annual median (A_MEDIAN; top-coded
            '#" rows set to the BLS top-code 239,200)
  jobzone   O*NET Job Zone (1-5, categorical)
  edu_cat   modal Required Level of Education category (O*NET 2.D.1 RL,
            1-12 scale, categorical)

WRITES occupation_market_join.json {soc: {...}} plus a coverage report.

Usage: python build_market_join.py
"""
import csv
import json
import os
import re
import statistics

import xlrd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(HERE, "market_anchors")


def parse_sps(path):
    s = open(path, encoding="latin-1").read()
    return {c: float(v) for c, v in
            re.findall(r"\(\s*(\d+)\s*=\s*(\d+(?:\.\d+)?)\)", s)}


def isco_score(table, isco4):
    """exact 4-digit, else the published aggregate (3/2/1-digit + zeros)."""
    for cand in (isco4, isco4[:3] + "0", isco4[:2] + "00", isco4[:1] + "000"):
        if cand in table:
            return table[cand]
        if cand.lstrip("0") in table:
            return table[cand.lstrip("0")]
    return None


def main():
    isei = parse_sps(os.path.join(A, "isqoisei08.sps"))
    siops = parse_sps(os.path.join(A, "isqotrei08.sps"))
    print(f"score tables: ISEI {len(isei)} codes, SIOPS {len(siops)}")

    wb = xlrd.open_workbook(os.path.join(A, "ISCO_SOC_Crosswalk.xls"))
    sh = next(s for s in wb.sheets() if s.name.startswith("2010 SOC"))
    soc2isco = {}
    for r in range(sh.nrows):
        vals = [str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)]
        m = re.match(r"^(\d{2}-\d{4})", vals[0])
        if m and re.match(r"^\d{4}$", vals[3].split(".")[0]):
            soc2isco.setdefault(m.group(1), []).append(vals[3].split(".")[0])
    print(f"crosswalk: {len(soc2isco)} SOC codes with ISCO-08 mappings")

    wagewb = load_workbook(os.path.join(A, "national_M2024_dl.xlsx"),
                           read_only=True)
    ws = wagewb.active
    head = [c.value for c in next(ws.iter_rows(max_row=1))]
    ic = head.index("OCC_CODE")
    im = head.index("A_MEDIAN")
    ig = head.index("O_GROUP") if "O_GROUP" in head else None
    wage = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if ig is not None and row[ig] not in ("detailed",):
            continue
        v = row[im]
        if v in ("*", None, ""):
            continue
        wage[str(row[ic]).strip()] = 239200.0 if v == "#" else float(v)
    print(f"OES wages: {len(wage)} detailed SOC codes")

    zone = {}
    for r in csv.DictReader(open(os.path.join(A, "job_zones.txt")),
                            delimiter="\t"):
        zone[r["O*NET-SOC Code"]] = int(float(r["Job Zone"]))
    edu = {}
    for r in csv.DictReader(open(os.path.join(A, "education_raw.txt")),
                            delimiter="\t"):
        if r["Element ID"] == "2.D.1" and r["Scale ID"] == "RL":
            k = r["O*NET-SOC Code"]
            v = float(r["Data Value"])
            if k not in edu or v > edu[k][1]:
                edu[k] = (int(r["Category"]), v)

    socs = sorted({r["soc"] for r in
                   json.load(open(os.path.join(HERE,
                                               "prereg_member_table.json")))
                   if r.get("soc")})
    out = {}
    miss = {"isei": 0, "wage": 0, "zone": 0, "edu": 0}
    for soc in socs:
        soc6 = soc[:7]
        iscos = soc2isco.get(soc6, [])
        es = [isco_score(isei, i) for i in iscos]
        ts = [isco_score(siops, i) for i in iscos]
        es = [x for x in es if x is not None]
        ts = [x for x in ts if x is not None]
        rec = {
            "isei": round(statistics.mean(es), 2) if es else None,
            "siops": round(statistics.mean(ts), 2) if ts else None,
            "isco_n": len(iscos),
            "wage": wage.get(soc6),
            "jobzone": zone.get(soc),
            "edu_cat": edu[soc][0] if soc in edu else None,
        }
        for k in ("isei", "wage"):
            if rec[k] is None:
                miss[k] += 1
        if rec["jobzone"] is None:
            miss["zone"] += 1
        if rec["edu_cat"] is None:
            miss["edu"] += 1
        out[soc] = rec
    json.dump(out, open(os.path.join(HERE, "occupation_market_join.json"),
                        "w"), indent=0)
    print(f"joined {len(out)} study SOC codes; missing — " +
          ", ".join(f"{k}: {v}" for k, v in miss.items()))
    multi = sum(1 for s in out.values() if s["isco_n"] > 1)
    print(f"SOC codes mapping to multiple ISCO-08 codes: {multi} "
          f"(mean of mapped scores taken)")


if __name__ == "__main__":
    main()
